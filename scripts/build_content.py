#!/usr/bin/env python3
"""Authoring source for the initial content + tuning data.

Emits:
  content/*.csv            (what the server loads)
  build/the-line-content.xlsx  (uploaded once to Google Sheets)

After the Google Sheet exists, the sheet is the source of truth. Re-sync with:
  npm run content:pull
"""
import csv
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTENT = os.path.join(ROOT, "content")
BUILD = os.path.join(ROOT, "build")

# ---------------------------------------------------------------------------
# Tuning
# ---------------------------------------------------------------------------
TUNING = [
    ("LINE_TICK_SECONDS", 60, "Seconds the front-of-liner has per interval. Also the line tick."),
    ("FRONT_EXTENSIONS_MAX", 2, "LOCKED. 'More time' presses allowed at the front."),
    ("EXIT_EMOJI_SECONDS", 20, "Seconds to pick an exit emoji after submitting before a default is used."),
    ("LEAVE_GRACE_SECONDS", 30, "Seconds without a heartbeat before a disconnected player loses their spot."),
    ("MESSAGE_MAX_CHARS", 140, "Max characters in a front-of-line message."),
    ("MESSAGE_DELIVERY_WEIGHT", 1.0, "Draw weight = 1 / (1 + deliveries) ^ this. Higher = stronger bias to undelivered messages."),
    ("SEED_RETIRE_THRESHOLD", 200, "Seeds stop being drawn once approved non-seed messages exceed this."),
    ("MOD_PASS_THRESHOLD", 0.7, "Classifier score >= this: auto-approve."),
    ("MOD_FAIL_THRESHOLD", 0.3, "Classifier score < this: auto-reject. Between the two: human review queue."),
    ("MESSAGE_REPORT_THRESHOLD", 3, "Reports needed to pull a message into the review queue."),
    ("XP_IDLE_PER_MINUTE", 2, "XP per full minute in line. Tuned low."),
    ("XP_FRONT_BONUS", 50, "XP granted the moment a player reaches the front. The hurrah."),
    ("STAT_BASE", 1, "Starting value of CHARM, INTELLIGENCE, STRENGTH."),
    ("QUEST_FIRST_DELAY_MINUTES", 1, "Minutes after joining before the first quest."),
    ("QUEST_INTERVAL_MINUTES", 4, "Minutes between quests."),
    ("QUEST_INTERVAL_JITTER_MINUTES", 1, "Plus/minus jitter on the quest interval."),
    ("QUEST_TIMEOUT_SECONDS", 180, "Unanswered quests expire silently after this."),
    ("QUEST_TIER_WEIGHT_BIAS", 2.0, "Tier draw weight = tier ^ this, over unlocked tiers. Higher = more top-tier quests."),
    ("QUEST_RECENT_MEMORY", 8, "Do not repeat any of the player's last N quests."),
    ("ROLL_BASE", 0.5, "ROLL_FORMULA: p = ROLL_BASE + (stat - difficulty) * ROLL_STEP, clamped."),
    ("ROLL_STEP", 0.08, "Probability change per point of (stat - difficulty)."),
    ("ROLL_FLOOR", 0.05, "Minimum success probability."),
    ("ROLL_CEILING", 0.95, "Maximum success probability."),
    ("QUEST_XP_FAIL", 5, "Consolation XP for a failed check. Never zero."),
    ("ANNOUNCE_LEVEL_MILESTONES", "5,10,15,20,25,30", "Level-ups that NPCs announce in chat."),
    ("ANNOUNCE_MIN_INTERVAL_SECONDS", 90, "Global minimum gap between NPC announcements. Missed events are dropped."),
    ("NPC_AMBIENT_INTERVAL_SECONDS", 45, "Average gap between NPC ambient chat lines (only while humans are in line)."),
    ("NPC_AMBIENT_JITTER_SECONDS", 30, "Plus/minus jitter on the ambient interval."),
    ("CHAT_HISTORY_LINES", 50, "Chat lines shown to a joining player."),
    ("CHAT_MAX_CHARS", 200, "Max characters per chat message."),
    ("CHAT_MOD_DELAY_SECONDS", 2, "Delay before a human chat line appears (real-time moderation window)."),
    ("CHAT_RATE_LIMIT_SECONDS", 3, "Minimum seconds between chat messages from one player."),
    ("LEADERBOARD_RESET_DAY", "Monday", "Weekly board resets at 00:00 UTC on this day."),
    ("LEADERBOARD_SIZE", 100, "Rows shown on each board."),
    ("NAME_MAX_CHARS", 20, "Max characters in a display name."),
    ("RECEIPT_SNAPSHOT_NAMES", 10, "How many names from the line appear on the text receipt."),
    ("LINE_VIEW_NAMES", 14, "How many names around the player are rendered in the line view."),
]

# Level -> cumulative XP required to reach it. Level 1 = 0.
XP_CURVE = []
total = 0
for lvl in range(1, 31):
    XP_CURVE.append((lvl, total))
    total += int(round(50 * (1.18 ** (lvl - 1))))

QUEST_TIERS = [
    # tier, unlock_level, difficulty_min, difficulty_max, xp_success
    (1, 1, 0, 2, 20),
    (2, 4, 2, 5, 40),
    (3, 8, 4, 8, 70),
    (4, 14, 7, 12, 110),
    (5, 20, 10, 18, 160),
]

ODDS_BANDS = [
    ("Long shot", 0.0),
    ("Risky", 0.2),
    ("Even", 0.4),
    ("Likely", 0.6),
    ("Sure thing", 0.8),
]

TITLES = [
    # stat, min_level, title, accessory (text avatar layer)
    ("NONE", 1, "Person in Line", ""),
    ("CHARM", 1, "Local Sweetheart", "sunglasses"),
    ("CHARM", 5, "Everyone's Cousin", "a very good haircut"),
    ("CHARM", 12, "Neighbourhood Fixture", "a scarf that people comment on"),
    ("CHARM", 20, "Mayor of This Block", "a sash, worn ironically, but not really"),
    ("INTELLIGENCE", 1, "Reads the Signs", "a clipboard"),
    ("INTELLIGENCE", 5, "Knows a Guy", "reading glasses pushed up"),
    ("INTELLIGENCE", 12, "Has a Spreadsheet", "a laminated map"),
    ("INTELLIGENCE", 20, "Line Theorist", "three pens in one pocket"),
    ("STRENGTH", 1, "Can Stand a While", "a thermos"),
    ("STRENGTH", 5, "Brought a Chair", "a folding chair under one arm"),
    ("STRENGTH", 12, "Load-Bearing", "a cooler, carried, never opened"),
    ("STRENGTH", 20, "Immovable Object", "a whole camping setup"),
]

AVATAR_PARTS = [
    ("build", "tall"), ("build", "short"), ("build", "average-height"), ("build", "lanky"), ("build", "sturdy"),
    ("hair", "buzz cut"), ("hair", "ponytail"), ("hair", "curtain bangs"), ("hair", "shaved head"), ("hair", "top bun"),
    ("hair", "mullet"), ("hair", "curly mop"), ("hair", "side part"), ("hair", "braids"), ("hair", "grey crop"),
    ("top", "grey hoodie"), ("top", "denim jacket"), ("top", "band tee"), ("top", "puffer vest"), ("top", "work polo"),
    ("top", "raincoat"), ("top", "cardigan"), ("top", "hi-vis jacket"), ("top", "oversized blazer"), ("top", "sports jersey"),
    ("bottom", "cargo shorts"), ("bottom", "black jeans"), ("bottom", "joggers"), ("bottom", "pleated skirt"),
    ("bottom", "chinos"), ("bottom", "bike shorts"), ("bottom", "corduroys"), ("bottom", "wide-leg trousers"),
    ("item", "tote bag"), ("item", "earbuds in"), ("item", "bucket hat"), ("item", "lanyard"), ("item", "backpack"),
    ("item", "nothing in particular"), ("item", "closed umbrella"), ("item", "coffee cup"), ("item", "phone on 3%"),
    ("item", "reusable bag with one thing in it"),
]

EMOJIS = [
    ("🎉", "\\o/", "party"),
    ("👋", "o/", "wave"),
    ("🙂", ":)", "smile"),
    ("😭", "T_T", "sobbing"),
    ("🤷", "shrug", "shrug"),
    ("🫡", "o7", "salute"),
    ("🚪", "[door]", "door"),
    ("💀", "x_x", "skull"),
    ("🙏", "_/\\_", "thanks"),
    ("🔥", "^^^", "fire"),
    ("🐦", "<(')", "pigeon"),
    ("🧍", "|o|", "standing"),
]

NPCS = [
    ("regular", "The Regular", "has been here since a product launch nobody remembers", "STRENGTH,INTELLIGENCE"),
    ("streamer", "The Streamer", "broadcasting the line to zero viewers", "CHARM,STRENGTH"),
    ("placeholder", "The Placeholder", "holding a spot for a friend who is obviously not coming", "CHARM,INTELLIGENCE"),
    ("vendor", "The Vendor", "sells one item, changes daily, never has change", "INTELLIGENCE,STRENGTH"),
    ("auditor", "The Auditor", "convinced they are in the wrong line, needs help proving it", "INTELLIGENCE,CHARM"),
]

NPC_LINES = [
    ("regular", "Nice day for it."),
    ("regular", "Line's moving. Slower than Tuesday, but moving."),
    ("regular", "Anyone seen the guy with the cooler? He was here in spring."),
    ("regular", "I used to count. Then I stopped. Then I started again. Now I just stand."),
    ("regular", "Somebody's getting to the front soon. Can feel it."),
    ("regular", "If you need the stool, ask. I only have the one."),
    ("regular", "This is about the length it was during the launch. Whichever launch."),
    ("regular", "There was a chalk mark here once. Right here. Meant something."),
    ("streamer", "Chat, we're at hour three. Hour three, chat."),
    ("streamer", "Shout out to everyone watching. Which is, at time of writing, nobody. Love you."),
    ("streamer", "Okay chat, somebody just joined the line. Say hi. Chat? Say hi."),
    ("streamer", "Battery's at 40. We're good. We're so good."),
    ("streamer", "If you're watching on replay: hi from the past. It was fine."),
    ("streamer", "New sub goal: one sub."),
    ("streamer", "I'm going to do a dance if we hit five viewers. We won't. But I would."),
    ("streamer", "Chat, real talk. This is the longest stream I've ever done. Chat?"),
    ("placeholder", "Still holding it. They're coming."),
    ("placeholder", "Sorry, that spot's taken. Not by me. By my friend."),
    ("placeholder", "They said they were leaving 'in a bit.' That was a bit ago."),
    ("placeholder", "If anyone's wondering about the gap in the line: that's on purpose."),
    ("placeholder", "My friend's going to be so glad I did this."),
    ("placeholder", "Do people usually text when they're on the way? Or just when they arrive?"),
    ("placeholder", "Three dots. Three dots. Gone. Okay."),
    ("placeholder", "I'm not leaving. That's the whole point of me."),
    ("vendor", "One item today. Ask me what it is."),
    ("vendor", "No change. I want to be clear about that up front."),
    ("vendor", "Tomorrow's item is different. I don't know what yet."),
    ("vendor", "If you're not buying, that's fine. Just don't touch."),
    ("vendor", "I've been asked for water. I don't sell water. Today."),
    ("vendor", "Somebody once offered me exact change. Best day of my life."),
    ("vendor", "The table's not for sale. People ask."),
    ("vendor", "Business is steady. Nobody buys anything. Steady."),
    ("auditor", "Has anyone here actually confirmed this is the right line?"),
    ("auditor", "There's a sign on the next street. This street has no sign. Just noting it."),
    ("auditor", "Current count: more than before. Filing that."),
    ("auditor", "If anyone has a map, I'd like to compare."),
    ("auditor", "The door has no handle on this side. I've checked twice."),
    ("auditor", "Who told you about this line? I'm collecting sources."),
    ("auditor", "I'm not saying it's the wrong line. I'm saying it's unconfirmed."),
    ("auditor", "Clipboard's for anyone. Just bring it back."),
]

ANNOUNCEMENTS = [
    ("regular", "long_shot", "Did everyone see that? {name} just pulled that off. I've been here years and I've never seen that."),
    ("streamer", "long_shot", "CHAT. {name} just did the thing. Clip it. Somebody clip it. Nobody's here. Clip it anyway."),
    ("placeholder", "long_shot", "{name} just did something incredible and I'm going to tell my friend about it when they get here."),
    ("vendor", "long_shot", "{name}. Long shot. Paid off. I'd sell that if I could."),
    ("auditor", "long_shot", "Noting for the record: {name} just beat the odds. Documented."),
    ("regular", "level_milestone", "{name} hit level {level}. That's a lot of standing. Respect."),
    ("streamer", "level_milestone", "Chat, {name} is level {level} now. Level {level}! We're all so proud. Chat? We're proud."),
    ("placeholder", "level_milestone", "{name}'s level {level}. I'd be that level too if I ever moved from this spot."),
    ("vendor", "level_milestone", "Level {level} for {name}. No discount. But well done."),
    ("auditor", "level_milestone", "{name}: level {level}. Verified. Congratulations. Filed."),
]

PROMPTS = [
    "Tell the next person one thing you're looking forward to.",
    "Describe the weather without using the word for it.",
    "What did you eat today?",
    "Leave a piece of advice you didn't take.",
    "Say something to someone who's tired.",
    "Describe the person in front of you without judging them.",
    "What's a small thing that went right today?",
    "Write what you'd want to read right now.",
    "Tell them something you noticed while waiting.",
    "Wish them something specific.",
    "What did you nearly do instead of this?",
    "Describe a sound you can hear right now.",
    "Tell them the time, in your own words.",
    "Leave a compliment for a stranger.",
    "What's the last thing that made you laugh?",
    "Tell them what your shoes are like.",
    "Say something true that isn't important.",
    "Recommend nothing. Just say hi.",
    "Describe your hands right now.",
    "Tell them one thing you'd tell yourself an hour ago.",
]

SEED_MESSAGES = [
    # (text, prompt index into PROMPTS)
    ("Whoever you are, I hope your shoes are comfortable. Mine were not.", 15),
    ("There was a pigeon earlier that looked at me like it knew something. Keep an eye out.", 8),
    ("I ate my emergency snack at minute nine. Learn from me.", 2),
    ("If the person behind you is humming, it's fine. They were behind me too.", 11),
    ("I don't know you but I'm proud of you for standing here.", 4),
    ("Somebody's phone alarm has been going off in a bag for twenty minutes. Nobody has claimed it.", 11),
    ("Check your pockets. You have more receipts than you think.", 3),
    ("I counted the windows on the building across the street. There are 41. Now you don't have to.", 8),
    ("My mum texted 'are you still in that line' and I said yes and she said 'ok' and that was the whole conversation.", 16),
    ("It's going to rain in a minute. Or it was, when I wrote this. Either way, hi.", 1),
    ("Don't lean on the post. It moves.", 3),
    ("I had a whole speech planned for this and now it's just: good luck with the rest of your day.", 9),
    ("Someone ahead of me was wearing two different shoes. On purpose, I think. Iconic.", 5),
    ("You've waited longer than you think. That's the thing about it.", 12),
    ("I'd tell you what I'm having for dinner but I don't know yet either.", 2),
    ("There's a bakery two streets over. Don't go. But it's there.", 10),
    ("I've been thinking about a song for an hour and I still don't know what it's called.", 16),
    ("Nobody's going to ask, so: how are you, actually?", 17),
    ("If you're reading this, someone stood exactly here and thought about you. That's all. That's the message.", 7),
    ("Standing still for this long makes you realise how much you usually move.", 8),
    ("I came here after work and now I'm not sure which one was the job.", 10),
    ("I hope the weather's the same as when I left. It was pretty good.", 1),
    ("The person in front of me was called Dev and he never said a word. Dev, if this is you, hi.", 5),
    ("I did not need to be here. I stayed anyway. That felt important to say.", 16),
    ("There's a vending machine on the corner and it only sells one thing and I still don't know what it is.", 8),
    ("My step count is huge. My steps were all in place.", 6),
    ("Nothing happened while I waited and it was the best afternoon I've had in weeks.", 6),
    ("Tell the next person something kind. I'm telling you: you're doing great.", 13),
    ("I met a stranger here who's now a person I'll think about occasionally forever.", 6),
    ("At one point everyone in the line looked up at the same time. I don't know why. Nothing was there.", 8),
    ("I made a list of things I'd do after this. I lost the list. I feel free.", 6),
    ("Somebody offered me half a sandwich. I said no and I regret it, and I'll regret it for years.", 3),
    ("The line moved twice while I wasn't looking and both times I felt like I'd missed something.", 19),
    ("I brought a book and read the same page four times. Good page.", 8),
    ("The awning drips. Stand slightly to the left.", 3),
    ("I'm writing this with cold hands. If it's warm where you are, enjoy it for me.", 18),
    ("The kid in front of me asked their dad what we were waiting for and he said 'you'll see' and I've been thinking about it since.", 14),
    ("Whatever you write next, make it better than this. It's not hard.", 7),
    ("I don't have advice. I have this: you're nearly there. That's a fact, not encouragement.", 4),
    ("My coffee's gone cold. It was cold when I got it. Something to think about.", 2),
    ("There's a sock on the pavement near the bins. It's been there the whole time. It might be there for you too.", 8),
    ("I texted 'nearly there' to three people and it was never true until now.", 12),
    ("If you see someone in a yellow raincoat, they were nice to me. Be nice to them.", 13),
    ("I wanted to say something profound. I'm hungry. That's what I've got.", 2),
    ("You'll forget most of this wait. I'm hoping you don't forget this bit.", 7),
    ("Somebody behind me is explaining the offside rule to nobody. It's been going on for a while.", 11),
    ("I got here early. It didn't matter. Nothing about getting here early mattered. I loved it anyway.", 19),
    ("Please don't tap the door. Someone did and we all felt it.", 3),
    ("I've been here so long my phone thinks this is my home.", 12),
    ("Hello from earlier. It's later now. That's the whole trick.", 17),
]

BLOCKLIST = [
    # Placeholder profanity list. Extend freely; matching is whole-word, case-insensitive,
    # with basic leetspeak normalisation (0->o, 1->i, 3->e, 4->a, 5->s, 7->t, @->a, $->s).
    "fuck", "fucking", "fucker", "motherfucker", "shit", "bullshit", "cunt", "bitch", "asshole",
    "dick", "cock", "pussy", "whore", "slut", "faggot", "fag", "retard", "retarded", "nigger", "nigga",
    "kike", "spic", "chink", "tranny", "rape", "rapist", "kys", "kill yourself",
]

# ---------------------------------------------------------------------------
# Quests. Columns:
# id, tier, npc_id, setup,
# a_label, a_stat, a_difficulty, a_success, a_failure,
# b_label, b_stat, b_difficulty, b_success, b_failure
# ---------------------------------------------------------------------------
Q = []


def quest(id_, tier, npc, setup, a, b):
    Q.append((id_, tier, npc, setup) + a + b)


C, I, S = "CHARM", "INTELLIGENCE", "STRENGTH"

# ---- Tier 1 ---------------------------------------------------------------
quest("regular_t1_a", 1, "regular",
      "The Regular squints at you. \"You're new. I can tell because you keep checking the time. I stopped doing that in, I want to say, March.\" They hold out a thermos. \"Coffee's gone cold. Still counts.\"",
      ("Drink the cold coffee in one go", S, 1,
       "You down it. It tastes like a filing cabinet. The Regular nods slowly, the way you'd nod at a dog that finally sat.",
       "You get halfway and something in the bottom shifts. The Regular takes the thermos back without comment, which is worse."),
      ("Ask how long they've been here", C, 1,
       "\"Since the thing with the phones,\" they say. You nod like you know. They seem to appreciate being nodded at.",
       "\"Long enough,\" they say, and then don't say anything else for a while. You have made it weird."))

quest("regular_t1_b", 1, "regular",
      "The Regular has a folded-up newspaper from a date you can't quite read. \"Crossword's half done. Four down. Seven letters. Clue is: patience.\"",
      ("Solve four down", I, 2,
       "\"WAITING,\" you say. They write it in, in pen, without checking. \"Huh,\" they say, and it's the warmest thing anyone's said to you today.",
       "\"STANDING,\" you say. Eight letters. The Regular looks at you, then at the paper, then back at you. \"It's fine,\" they say. It is not fine."),
      ("Hold the paper flat while they write", S, 0,
       "You hold it steady against the wind. They finish the clue themselves. Teamwork, technically.",
       "A gust takes the sports section. You watch it go. The Regular watches you watch it go."))

quest("streamer_t1_a", 1, "streamer",
      "The Streamer swings the phone toward you. \"Chat, we've got a new face. Say hi to chat.\" The screen says LIVE · 0.",
      ("Say hi to chat", C, 1,
       "\"Hi chat,\" you say, and give a little wave. The Streamer says \"Chat loves you\" with total sincerity. Nobody is watching. It still feels nice.",
       "\"Hey... chat,\" you say, to the number zero. The Streamer says \"Awkward!\" in a bright voice and moves the camera away."),
      ("Hold the phone so they can fix their hair", S, 1,
       "You hold it steady at the exact angle they wanted. They say \"You're a natural, honestly.\" Zero people see this.",
       "You hold it at slightly the wrong angle. They take it back. \"It's fine. It's fine. Chat, it's fine.\""))

quest("streamer_t1_b", 1, "streamer",
      "\"Okay chat, sponsor segment.\" The Streamer turns to you. \"I don't have a sponsor. Can you pretend to be one for like ten seconds?\"",
      ("Improvise a sponsor read", C, 2,
       "You recommend a fictional insole brand with such conviction the Streamer asks for a discount code. You make one up. They write it down.",
       "You start strong and then say \"and it's, um, waterproof\" about a product you have not named. The Streamer says \"We'll cut that.\" There is no editor."),
      ("Explain that fake sponsorships are a disclosure issue", I, 1,
       "You cite a regulation you are fairly sure exists. The Streamer says \"Chat, we're being compliant today\" and seems genuinely relieved.",
       "You start explaining and lose the thread halfway. The Streamer nods along and says \"Legal stuff, chat.\" You have not helped."))

quest("placeholder_t1_a", 1, "placeholder",
      "The Placeholder is holding a spot with both arms out, like the spot might escape. \"My friend's coming. Can you just, like, confirm to people that this spot is taken? If they ask.\"",
      ("Vouch for the friend", C, 1,
       "You tell the person behind that the spot's taken. They shrug. The Placeholder mouths \"thank you\" with tears nearly in their eyes.",
       "You say \"their friend is coming\" in a tone that makes it clear you don't believe it. The Placeholder heard the tone."),
      ("Ask what time the friend said", I, 1,
       "\"Ten minutes ago,\" they say. You point out that ten minutes ago is not a time. They laugh, a little, for the first time today.",
       "\"Soon,\" they say. You ask \"soon like when?\" and they stop making eye contact. You have touched something."))

quest("placeholder_t1_b", 1, "placeholder",
      "\"Can you watch my spot AND my friend's spot?\" The Placeholder holds up two fingers. \"I need to check if they texted. I'll be one second.\"",
      ("Stand wide enough for two people", S, 1,
       "You occupy two spots with the stance of someone who has done this before. Nobody challenges you. The Placeholder returns. No text.",
       "Someone drifts into the second spot. You don't have it in you to stop them. The Placeholder returns and does the maths."),
      ("Talk the people behind into respecting the gap", C, 2,
       "You explain the situation with warmth. They agree it's a lot. Everyone stands respectfully around an empty rectangle of pavement.",
       "You explain the situation and someone says \"that's not really how lines work.\" They are right. You have no counter."))

quest("vendor_t1_a", 1, "vendor",
      "The Vendor has a folding table and one item on it: a single umbrella. It is not raining. \"Four fifty,\" they say. \"No change.\"",
      ("Point out it's not raining", I, 1,
       "\"Not yet,\" they say, and glance at a completely clear sky with such confidence you check it too.",
       "\"Not yet,\" they say. You say \"well the forecast is\" and they say \"the forecast is not here.\" You have lost."),
      ("Help move the table two feet left", S, 1,
       "The table is heavier than it looks. You move it. The Vendor says \"better\" without explaining what's better.",
       "You lift one end and the umbrella slides off. The Vendor picks it up and puts it back, exactly where it was, and says nothing."))

quest("vendor_t1_b", 1, "vendor",
      "\"Today's item,\" says the Vendor, gesturing at a single AA battery. \"Just the one. Two dollars. I don't have change.\"",
      ("Ask what it's for", I, 2,
       "\"Depends on you,\" they say. You consider this longer than you'd like to admit.",
       "\"Batteries,\" they say. You say \"yes but for\" and they say \"batteries.\" You buy nothing. You learn nothing."),
      ("Compliment the merchandising", C, 1,
       "\"Thank you,\" they say. \"It's about restraint.\" They straighten the battery by a millimetre.",
       "\"Nice setup,\" you say. \"It's a table,\" they say. You agree. It is a table."))

quest("auditor_t1_a", 1, "auditor",
      "The Auditor has a clipboard. \"Quick one. Is this the line for the thing? Because I was told the line for the thing was on the other side.\"",
      ("Assure them this is the line", I, 1,
       "\"This is the line,\" you say. They look at the clipboard, then at you, and write something down. \"Okay. Okay.\" They stay.",
       "\"This is the line,\" you say. \"For what?\" they say. You have no answer. They write that down too."),
      ("Say you've been wondering the same thing", C, 1,
       "You bond over uncertainty. They show you a diagram of the block. It's quite good, actually.",
       "You say you've been wondering too and they say \"well SOMEBODY should know\" loud enough that people turn around."))

quest("auditor_t1_b", 1, "auditor",
      "\"I've counted,\" says the Auditor, \"and there are more people here than there were an hour ago, which means the line is getting LONGER, which is not how lines are supposed to go.\"",
      ("Explain that people join at the back", I, 2,
       "You explain arrivals versus departures. They nod slowly. \"A flow problem,\" they say, and draw an arrow on the clipboard.",
       "You start explaining and use the phrase \"net throughput.\" They stare at you. \"So it IS getting longer,\" they say."),
      ("Stand still and be counted", S, 0,
       "You stand very still while they count. \"Thirty-two,\" they say. You don't know if that's good. Neither do they.",
       "You shift your weight and they lose count. \"From the top,\" they say, and they mean it."))

# ---- Tier 2 ---------------------------------------------------------------
quest("regular_t2_a", 2, "regular",
      "The Regular produces a small folding stool from nowhere. \"You can have the stool for a bit. But you have to tell me something you've never told anyone in this line.\"",
      ("Tell them something true", C, 3,
       "You tell them. They listen the whole way through. \"Yeah,\" they say. \"Stool's yours till the line moves.\" It's a good stool.",
       "You start with \"honestly\" and then say something you've definitely told people. They know. The stool goes back into wherever it came from."),
      ("Decline the stool. Standing is the point.", S, 3,
       "\"Standing's the point,\" you say. The Regular looks at you for a long time. \"Yeah,\" they say finally. \"Yeah it is.\"",
       "You decline the stool and then your knee does something. The Regular sees. You both pretend they didn't."))

quest("regular_t2_b", 2, "regular",
      "\"They used to hand out water,\" says the Regular. \"Little cups. There's a box behind that post that might still have some.\"",
      ("Go get the box without losing your spot", S, 4,
       "You lunge, grab, and return in one motion. The box has eleven paper cups and no water. The Regular seems thrilled anyway.",
       "You get to the post and somebody says \"you're out of line\" in a tone that has layers. You come back empty-handed and slightly rattled."),
      ("Ask which post, exactly", I, 3,
       "You ask enough clarifying questions to establish that the box is real, was real, and is now probably not. The Regular respects the process.",
       "\"That post,\" they say, pointing at four posts. You pick one. It's the wrong one. It was always going to be."))

quest("streamer_t2_a", 2, "streamer",
      "\"Chat wants a Q&A,\" the Streamer says, checking a chat with nothing in it. \"You're the guest. First question: what's your content?\"",
      ("Describe yourself as a brand", C, 4,
       "You pitch yourself as \"lifestyle, but standing.\" The Streamer says \"That's a niche\" and follows you on a platform you're not on.",
       "You say \"I don't really have content\" and the Streamer says \"That's okay!\" in the voice of someone who thinks it isn't."),
      ("Ask about their analytics", I, 2,
       "You ask about watch time. They pull up a dashboard that is one flat line. \"Consistent,\" you say. They love that.",
       "You ask about analytics and they show you a graph with no y-axis. You ask about the y-axis. It's not a good moment."))

quest("streamer_t2_b", 2, "streamer",
      "\"Okay chat, we're doing a bit.\" The Streamer hands you a second phone. \"You're the rival streamer. We're beefing. Say something cutting.\"",
      ("Deliver a devastating fake insult", C, 3,
       "You say \"At least MY chat has a zero in it too.\" The Streamer gasps. \"Chat, clip that.\" Nobody clips it. It lives in your heart.",
       "You say \"your... shoes are bad.\" The Streamer looks at their shoes. They're fine shoes. The bit dies."),
      ("Hold both phones steady for the two-camera setup", S, 3,
       "You become a tripod. Both angles are clean. The Streamer says \"production value\" with tears in their eyes.",
       "Your arm gives out on the second phone. The camera captures pavement. \"B-roll,\" the Streamer says weakly."))

quest("placeholder_t2_a", 2, "placeholder",
      "The Placeholder's phone buzzes. They look at it and their whole face changes. \"They're... running late. Later than before. Can you help me figure out what to say back?\"",
      ("Draft a reply that gets a real answer", I, 3,
       "You write: \"No worries, what time should I stop holding it?\" The Placeholder stares at it. \"That's... direct.\" They send it. Three dots appear. The dots go away.",
       "You write \"lol no rush\" and the Placeholder sends it and there is, indeed, no rush. There will never be rush."),
      ("Suggest, gently, that they might not come", C, 4,
       "You say it softly. They nod. \"I know,\" they say. \"But I already told the guy behind us.\" They keep holding the spot, but differently now.",
       "You say it and they say \"They're COMING\" and turn away, and you have somehow made a stranger's day worse."))

quest("placeholder_t2_b", 2, "placeholder",
      "\"Okay here's the thing,\" says the Placeholder. \"I need the bathroom. Can you hold MY spot and THEIR spot and if anyone asks, I'm right back.\"",
      ("Hold two spots with your body", S, 3,
       "You spread out like a barricade. Nobody passes. The Placeholder returns, relieved in every sense.",
       "You spread out but the line moves and you can't spread forward. The Placeholder returns to one spot. Their friend's spot is a memory now."),
      ("Explain to the neighbours", C, 2,
       "You give a short, sincere briefing. Everyone agrees to honour the gap. It's genuinely moving.",
       "You explain and someone says \"wait, so the friend isn't even here?\" and the whole arrangement collapses under scrutiny."))

quest("vendor_t2_a", 2, "vendor",
      "The Vendor's item today is a single left glove. \"Good glove. Eight dollars. No change. And before you ask: I don't know where the right one is.\"",
      ("Negotiate on the grounds that it's half a product", I, 4,
       "You argue for a fifty percent discount. The Vendor considers this seriously. \"Four dollars,\" they say. \"No change.\" You have won and also can't pay.",
       "You argue it's half a product. \"It's a whole glove,\" they say. You can't argue with that. You try. You can't."),
      ("Try the glove on to prove it fits", S, 3,
       "It fits. It fits really well. You don't buy it, but you think about it for the rest of the day.",
       "It's tight. You get it on and can't get it off. The Vendor watches. Eventually it comes off. \"Eight dollars,\" they say."))

quest("vendor_t2_b", 2, "vendor",
      "\"Today's item,\" says the Vendor, unveiling a single houseplant. \"Needs water. Six dollars. No change. If you're not buying, it still needs water.\"",
      ("Find water for the plant", S, 2,
       "You produce a water bottle and give the plant a drink. The Vendor says \"It likes you,\" which you didn't know plants could do.",
       "You go looking for water, get caught in a conversation you weren't in, and come back with nothing. The plant looks worse. That's probably not your fault."),
      ("Identify the plant", I, 5,
       "\"Pothos,\" you say. The Vendor looks impressed for the first time in recorded history. \"Seven dollars now,\" they say.",
       "\"Fern,\" you say. It's not a fern. The Vendor says \"Fern\" back to you in a way that lets you know."))

quest("auditor_t2_a", 2, "auditor",
      "The Auditor is holding up a photo of a different line. \"Look. THAT line has a sign. This line has no sign. Ergo.\"",
      ("Explain why absence of a sign proves nothing", I, 3,
       "You walk them through it. They write \"ABSENCE IS NOT EVIDENCE\" on the clipboard and underline it twice. They stay.",
       "You say \"signs aren't everything\" and they say \"signs are LITERALLY everything\" and then show you the photo again."),
      ("Ask to see the photo properly", C, 3,
       "You look at it with real interest. It's a good photo. You say so. They soften. \"It IS a good photo.\" The line thing is forgotten for now.",
       "You look and say \"huh.\" They wait for more. There is no more. They put the photo away like you've hurt it."))

quest("auditor_t2_b", 2, "auditor",
      "\"I've drawn up a form,\" says the Auditor. \"For confirming you're in the right line. Could you fill it in? It's for my records.\"",
      ("Fill in the form accurately", I, 4,
       "Question 3 is \"Purpose of visit.\" You leave it blank. The Auditor looks at the blank, then at the door, then nods. \"Fair.\"",
       "You put \"N/A\" for every field. The Auditor says \"you can't N/A your own name\" and they're right."),
      ("Hold the clipboard steady while they fill it in", S, 2,
       "You hold it at exactly the right height. They write with great satisfaction. The form is about you. You don't ask.",
       "The pen goes through the paper. The Auditor sighs and starts a fresh copy. They have many copies."))

# ---- Tier 3 ---------------------------------------------------------------
quest("regular_t3_a", 3, "regular",
      "The Regular leans in. \"Between us. There was a day, a while back, where the line moved backwards. Just once. Nobody believes me.\"",
      ("Work out how that could have happened", I, 6,
       "You reason it through: a merge, a miscount, someone rejoining at the wrong spot. The Regular listens, eyes shining. \"So you believe me.\" You sort of do.",
       "You say \"lines can't move backwards\" and they say \"that's what THEY said\" and you realise you've become one of THEY."),
      ("Just believe them", C, 5,
       "\"I believe you,\" you say. They exhale like they've been holding it since that day. \"Thank you.\" You feel it too, somehow.",
       "You say \"sure\" in a way that has a shrug in it. They hear the shrug. They always hear the shrug."))

quest("regular_t3_b", 3, "regular",
      "\"Someone once got to the front,\" the Regular says, \"and turned around and got back in line. Right at the back. I think about that a lot. What would you do?\"",
      ("Say you'd do it too, and mean it", S, 6,
       "\"I'd get back in,\" you say. The Regular studies you. \"Yeah. You would.\" It's the highest compliment they have.",
       "You say you'd get back in, and they ask \"why,\" and you don't have a why, and it shows."),
      ("Ask what that person's face looked like", I, 5,
       "The Regular thinks for a while. \"Calm,\" they say. \"Real calm.\" You both look at the door. The door does nothing.",
       "You ask and they say \"I don't remember faces, I remember positions,\" and then they describe the position in detail. It's a lot."))

quest("streamer_t3_a", 3, "streamer",
      "\"Chat, huge news.\" The Streamer is shaking. \"We just got a viewer. One. And they're asking who YOU are.\" The screen says LIVE · 1.",
      ("Introduce yourself to the one viewer", C, 6,
       "You look into the lens and say your name and where you are. The viewer types \"cool.\" The Streamer nearly cries. It's the best day of their career.",
       "You say \"uh, hi\" and the viewer leaves. LIVE · 0. The Streamer says \"It's fine, they had to go, chat, they had to go.\""),
      ("Check whether the viewer is a bot", I, 5,
       "You ask the viewer a question a bot couldn't answer. They answer it. Real person. The Streamer's hands are shaking.",
       "You ask \"are you a bot\" and the viewer types \"no\" and leaves. The Streamer says \"You scared them.\" You did."))

quest("streamer_t3_b", 3, "streamer",
      "The Streamer's battery is at four percent. \"I have to keep the stream up. It's the longest one I've ever done. Do you have... anything?\"",
      ("Sprint to a nearby outlet and back before your spot's gone", S, 5,
       "You go, you plug in for ninety seconds, you return. Twelve percent. The stream lives. The Streamer calls you \"the community.\"",
       "You go, and there's a queue for the outlet. Of course there is. You come back to a black screen and a Streamer staring at their reflection in it."),
      ("Talk them through low-power settings", I, 7,
       "Brightness down, background apps off, resolution dropped. Four percent becomes an hour. \"You're a wizard,\" they say, and they mean it.",
       "You suggest a setting that doesn't exist on their phone. They tap around looking for it. Three percent. Two."))

quest("placeholder_t3_a", 3, "placeholder",
      "The Placeholder is very quiet. \"They texted. They said they're not coming. They said I should just go in without them.\" They don't move. \"I don't know what I'm holding the spot for now.\"",
      ("Tell them the spot was always theirs", C, 6,
       "\"It was your spot,\" you say. \"It was always your spot.\" They nod, slowly, and step into it properly for the first time.",
       "You say \"well now you can just wait normally\" and they say \"normally\" like it's a word in a language they don't speak."),
      ("Point out they've still waited the whole time", I, 6,
       "You explain that the wait counted whether or not the friend showed. They think about it. \"It counted,\" they repeat. Something settles.",
       "You explain the wait still counted and they say \"counted toward WHAT\" and you look at the door and neither of you says anything."))

quest("placeholder_t3_b", 3, "placeholder",
      "\"New plan,\" says the Placeholder, brighter than they've been all day. \"I'm going to hold a spot for the NEXT person who needs one. Anyone. Can you help me find someone?\"",
      ("Find someone in line who needs a spot held", C, 5,
       "You find someone who needs to step out for a minute. The Placeholder holds their spot with visible joy. They have a purpose again.",
       "You ask around and everyone says they're fine. The Placeholder holds a spot for no one, which, now that you think about it, they've been doing all along."),
      ("Suggest they hold the spot for themselves, formally", I, 4,
       "You draft a short declaration. They sign it. They are now, on paper, holding a spot for themselves. It's oddly moving.",
       "You suggest it and they say \"that's not how holding works\" and you don't have the energy to argue about holding."))

quest("vendor_t3_a", 3, "vendor",
      "The Vendor's item today is a small locked box. \"Twelve dollars. No key. No change. I don't know what's in it. That's not a sales pitch, that's just true.\"",
      ("Deduce the contents from the weight and sound", I, 7,
       "You shake it, tilt it, listen. \"Coins,\" you say. \"About four dollars' worth.\" The Vendor's eyes narrow. \"Eight dollars, then.\"",
       "You shake it. Something rattles. \"It's... something,\" you say. The Vendor says \"Twelve dollars for something.\" You have not narrowed it down."),
      ("Offer to open it by force", S, 6,
       "You don't open it, but you make it clear you could. The Vendor moves it slightly further from you. \"Ten dollars,\" they say. Respect.",
       "You try to pry it and the lid doesn't move but your thumbnail does. The Vendor offers you a plaster for a dollar. No change."))

quest("vendor_t3_b", 3, "vendor",
      "\"Today,\" says the Vendor, \"the item is a receipt. For something. Fifteen dollars. It's a good receipt. Thermal paper. No change.\"",
      ("Read the receipt for clues", I, 6,
       "The store name is faded. The total is $15.00. The date is today. The Vendor is selling you a receipt for the receipt. You feel like you've solved something.",
       "You squint. The text is all faded. \"It was clearer this morning,\" says the Vendor. You believe them, somehow."),
      ("Ask them why they sell one thing a day", C, 5,
       "\"One thing is enough,\" they say. \"More than one and people start choosing.\" It's the most anyone's explained anything to you today.",
       "\"Why one thing?\" you ask. \"Why not?\" they say. You have no follow-up. You should have prepared a follow-up."))

quest("auditor_t3_a", 3, "auditor",
      "The Auditor has a laminated map. \"I've cross-referenced every line in a six-block radius. This one isn't on any list. Which means either it doesn't exist, or it's new. Which is it?\"",
      ("Argue it's simply unlisted", I, 7,
       "You explain that lists are lagging indicators. They look at the map, then at the very real line they're standing in. \"Unlisted,\" they whisper, and laminate a note.",
       "You say \"maybe it's a secret\" and the Auditor's eyes go wide in a way you'll regret for the next twenty minutes."),
      ("Compliment the map", C, 6,
       "\"This is really thorough,\" you say. They light up. They walk you through the legend. There's a symbol for benches. You forget about the question. So do they.",
       "\"Nice map,\" you say, and they say \"it's not NICE, it's ACCURATE\" and you have picked the wrong adjective."))

quest("auditor_t3_b", 3, "auditor",
      "\"Okay,\" says the Auditor. \"I'm going to go ask someone at the front. If I'm in the wrong line I need to know. Can you hold my spot while I go?\"",
      ("Hold the spot against all comers", S, 7,
       "You hold it. Two people try to close the gap. You don't budge. The Auditor returns with no information and their spot intact. \"Nobody at the front knows anything,\" they say, delighted.",
       "The line moves and you get carried forward. The Auditor returns to find their spot has become a concept."),
      ("Convince them the front won't know either", I, 5,
       "You explain that the front is just a person who waited longer. They think about that. \"Then nobody knows,\" they say, and for once that seems to calm them.",
       "You say \"the front won't know\" and they say \"then WHO\" with such volume that the front turns around."))

# ---- Tier 4 ---------------------------------------------------------------
quest("regular_t4_a", 4, "regular",
      "The Regular looks tired in a new way. \"I've been thinking about leaving. Not the line. Just... going home for a night. Sleeping in a bed. Would you think less of me?\"",
      ("Tell them the line will still be here", C, 9,
       "You say it warmly. They laugh, a real one. \"Yeah. It's not going anywhere.\" They don't leave. But they could now, and that's different.",
       "You say \"go home, honestly\" and they hear \"you're done here\" and that is not what you meant at all."),
      ("Offer to hold their spot overnight", S, 8,
       "\"I'll hold it,\" you say. \"Whole night.\" They look at you like you've offered a kidney. They don't take it. They needed to know it was there.",
       "You offer, and then yawn mid-sentence. The Regular pats your arm. \"You're sweet. You wouldn't make it.\""))

quest("regular_t4_b", 4, "regular",
      "\"You know what I've never done,\" the Regular says. \"I've never been at the front. Every time I get close, I let it go. I don't think I want to know.\" They look at you. \"Do you?\"",
      ("Give an honest answer about wanting to know", I, 10,
       "You think about it properly, and say what's true. The Regular listens. \"Huh,\" they say. \"That's more than I've worked out.\" Something shifts between you.",
       "You say something clever about the journey and the destination and the Regular says \"that's a poster\" and you deserve that."),
      ("Ask them to go to the front with you, when it's time", C, 8,
       "They're quiet for a long moment. \"Maybe,\" they say. It's the first maybe they've said in years.",
       "You ask and they shake their head hard. \"No. No. That's not how it works.\" You've pushed too far."))

quest("streamer_t4_a", 4, "streamer",
      "The Streamer's phone shows LIVE · 3. They're vibrating. \"Three. THREE. One of them is asking what the line is for. What do I say? What do I SAY?\"",
      ("Take the phone and answer for them", C, 10,
       "\"It's the line,\" you say, into the camera, with total calm. Chat: \"ok.\" \"lol.\" \"fair.\" Three viewers stay. The Streamer is holding your arm.",
       "You take the phone and say \"so basically\" and then you're still saying \"basically\" twenty seconds later and viewers drop to one."),
      ("Coach them on what not to say", I, 8,
       "You tell them: don't explain, don't guess. Just be here. They nod, turn to the camera, and say \"We're here, chat.\" Chat approves. Chat is three people.",
       "You tell them to \"stay mysterious\" and they say \"it's a MYSTERY, chat\" in a spooky voice and a viewer leaves. Then another."))

quest("streamer_t4_b", 4, "streamer",
      "\"I'm going to do a 24-hour stream,\" the Streamer says. \"Of the line. I need a co-host for the night shift. You in?\"",
      ("Commit to the night shift", S, 11,
       "You say yes and you mean it. Hour nine, you're reading chat aloud to no one, and the Streamer is asleep on your shoulder. Great content.",
       "You say yes and are asleep by hour two. The stream captures you snoring at 720p. It gets the highest view count they've ever had."),
      ("Convince them 24 hours is a bad idea", C, 9,
       "You make the case for quality over quantity. They listen. \"Six hours,\" they say. \"Six hours and a highlight reel.\" Growth.",
       "You say it's a bad idea and they say \"you sound like my sister\" and now you're their sister, apparently."))

quest("placeholder_t4_a", 4, "placeholder",
      "The Placeholder is holding a spot for their friend again. \"They said they're coming. This time for real. They sent a screenshot of their route.\" They show you. The route is from the airport.",
      ("Estimate the arrival time from the route", I, 9,
       "Two hours, you say, allowing for traffic. The Placeholder does some maths, then some more. \"Two hours is fine,\" they say, and for the first time you believe the friend is real.",
       "You say \"forty minutes\" and are wrong by a lot. When the friend doesn't appear at forty minutes, the Placeholder looks at you like you promised."),
      ("Offer to help hold the spot until then", C, 8,
       "You take turns leaning into the gap. It becomes a shift system. A stranger joins. The friend arrives. The friend is real. Everyone cries a little.",
       "You offer, and then you drift. The gap closes. The friend arrives to a Placeholder holding air."))

quest("placeholder_t4_b", 4, "placeholder",
      "\"I've been holding this spot so long,\" the Placeholder says, \"I don't think I know how to just... be in line. As myself. Does that make sense?\"",
      ("Help them practise being in line as themselves", C, 10,
       "You stand together, not holding anything. Just standing. \"This is weird,\" they say. \"Good weird.\" They stop looking at their phone.",
       "You say \"just stand there\" and they stand there so rigidly that people ask if they're okay. They are not, quite."),
      ("Reframe the holding as a skill", I, 8,
       "You point out that holding space for someone is a thing people pay for. They laugh. \"A professional.\" The word suits them.",
       "You say \"holding is a skill\" and they say \"a skill for WHAT\" and you say \"holding\" and the loop does not close."))

quest("vendor_t4_a", 4, "vendor",
      "The Vendor's item today is a folding chair. \"Twenty dollars. No change. It's a good chair. Sitting is not permitted in the line, but I don't make the rules. I just sell the chair.\"",
      ("Carry the chair, unfolded, for the rest of your wait, without sitting", S, 10,
       "You carry it like a shield. People look. You never sit. The Vendor says \"I've never seen anyone do that\" and gives you a discount you can't use.",
       "You carry it for a while and then your arms make a decision without you. You sit. Briefly. Everyone saw."),
      ("Argue the chair is a lifestyle product, not a seat", I, 9,
       "You make a case involving posture and intention. The Vendor listens. \"A lifestyle chair,\" they say, and change the price tag. Up.",
       "You argue and the Vendor says \"it's a chair, you sit on it\" and folds it, unfolds it, and folds it again as a demonstration."))

quest("vendor_t4_b", 4, "vendor",
      "\"Today's item,\" says the Vendor, \"is a number. Seventeen. Twenty-five dollars. It's a good number. No change.\" There is nothing on the table.",
      ("Buy the number, conceptually", I, 11,
       "You ask what you get. \"Seventeen,\" they say. You shake on it. You now own seventeen. You feel, strangely, wealthier.",
       "You ask what you get and they say \"seventeen\" and you say \"but what IS it\" and they say \"it's seventeen\" and this goes on."),
      ("Convince them to throw in eighteen", C, 9,
       "\"Two for one,\" you say, and smile. The Vendor considers. \"Eighteen's not for sale.\" Pause. \"But I'll mention you had a good attitude.\"",
       "You ask for eighteen and they say \"eighteen is TOMORROW'S item\" with such offence that you apologise to a number."))

quest("auditor_t4_a", 4, "auditor",
      "The Auditor sits down on the kerb, clipboard in lap. \"I've been in a lot of lines. Wrong ones, mostly. And I'm starting to think I might be in the right one this time, and I don't know what to do with that.\"",
      ("Sit with them", C, 9,
       "You sit. Neither of you says anything for a while. \"This is fine,\" they say eventually. \"Being right is fine.\" They don't get up for a long time.",
       "You sit and immediately start reassuring them and they say \"I didn't ask for that\" and they're right, they didn't."),
      ("Help them audit this line properly", I, 10,
       "You go through the checklist together. Location, duration, direction of travel. Every box gets a tick. \"Right line,\" they say. \"Right line.\" They frame the page.",
       "You go through the checklist and get to \"purpose\" and neither of you can fill it in and the audit fails on a technicality."))

quest("auditor_t4_b", 4, "auditor",
      "\"Someone told me,\" says the Auditor, \"that there's a person at the front who checks names. I've never seen them. I've never seen anyone come OUT. Have you?\"",
      ("Reason carefully about where people go", I, 12,
       "You lay out the possibilities without committing to any. The Auditor listens, writes, nods. \"So we don't know,\" they say. \"Officially.\" Officially seems to help.",
       "You say \"they probably come out the back\" and the Auditor says \"there IS no back, I've checked\" and shows you a diagram of the back."),
      ("Walk to the door and look, then walk back", S, 8,
       "You go. You look. It's a door. You come back. \"Door,\" you report. The Auditor writes \"DOOR (CONFIRMED)\" and seems satisfied.",
       "You go and someone at the front says \"excuse me\" in a way that makes you retreat immediately. \"Well?\" says the Auditor. You have nothing."))

# ---- Tier 5 ---------------------------------------------------------------
quest("regular_t5_a", 5, "regular",
      "The Regular is three from the front. They haven't been this close in years. Their hands are shaking. \"I'm going to step out,\" they say. \"I always step out. Tell me not to.\"",
      ("Tell them not to", C, 14,
       "You tell them. Simply. They stay. Two from the front now. One. They look back at you once, and then they go. You don't know what happened after.",
       "You say \"stay\" and they say \"why\" and you don't have the why and they step out and rejoin at the back. Same as always."),
      ("Stand behind them so there's no room to step out", S, 13,
       "You close the gap with your whole body. They feel it. \"Oh,\" they say. They can't step back. So they step forward.",
       "You close the gap and they turn sideways and slip out anyway. \"Sorry,\" they say, already walking to the back."))

quest("regular_t5_b", 5, "regular",
      "\"I want to give you something,\" the Regular says. It's the thermos. \"It's been in this line longer than anyone. I think it should go on without me.\"",
      ("Accept it, and carry it the rest of the way", S, 12,
       "It's heavier than it looks. Full of something. You don't check. You carry it. When you get to the front, it's still with you. It's yours now.",
       "You accept it and set it down for a second and it's gone. Just gone. The Regular doesn't say anything, which is somehow the worst thing."),
      ("Ask what's in it", I, 11,
       "\"Cold coffee,\" they say. \"From the first day.\" You look at them. They're not joking. You take it anyway. It's history.",
       "You ask and they say \"don't\" and you ask again and they take it back. Some things you're supposed to just carry."))

quest("streamer_t5_a", 5, "streamer",
      "The Streamer's phone reads LIVE · 40. Forty. Chat is moving too fast to read. \"They want to know what happens at the front,\" they say, white as a sheet. \"They keep asking. What do I do?\"",
      ("Tell chat, on camera, that they'll find out when they line up", C, 15,
       "You say it with a grin. Chat explodes. \"lmao.\" \"ok fine.\" \"brb joining.\" The count hits fifty. The Streamer sits down on the pavement.",
       "You say it and it comes out smug. Chat turns. \"gatekeeping.\" \"ratio.\" Thirty. Twenty. The Streamer takes the phone back gently."),
      ("Moderate the chat", I, 13,
       "You set slow mode, pin a message that says \"we don't know either,\" and time out two bots. Chat calms. The Streamer looks at you like you're a paramedic.",
       "You try to set slow mode and accidentally end the stream. LIVE · 0. The Streamer stares at the black screen. \"It's fine. It's fine. We'll go again.\""))

quest("streamer_t5_b", 5, "streamer",
      "\"They offered me a sponsorship,\" the Streamer whispers. \"A real one. Insoles. But they want me to say what the line is for. I don't know what the line is for. What do I do?\"",
      ("Draft a sponsor read that says nothing", I, 14,
       "You write: \"Whatever's at the front, you'll want to be comfortable getting there.\" The Streamer reads it. The sponsor loves it. Nobody said anything.",
       "You write something that accidentally implies the line is for insoles. The sponsor loves it. The Streamer doesn't. It goes out anyway."),
      ("Convince them to turn it down", C, 13,
       "You say \"you don't need it.\" They look at their forty viewers. \"I don't need it,\" they repeat. They decline. Chat cheers. It's the best content they've ever made.",
       "You say turn it down and they say \"easy for you to say\" and they're right, it is easy for you to say."))

quest("placeholder_t5_a", 5, "placeholder",
      "The friend arrived. They're here. Real. They're standing in the spot. The Placeholder is standing just behind, looking at their own hands. \"I don't know what to do now,\" they say quietly.",
      ("Tell the friend what the Placeholder did", C, 14,
       "You tell them. All of it. The hours. The gap. The friend turns around and looks at the Placeholder like they're seeing them for the first time. They swap places.",
       "You tell the friend and the friend says \"yeah, I know, thanks\" without turning around and the Placeholder says \"it's fine\" and it is not."),
      ("Point out the Placeholder is now, technically, behind", I, 12,
       "You explain the geometry. The friend gets it immediately, steps back, and pushes the Placeholder forward. \"You waited,\" they say. \"You go first.\"",
       "You explain the geometry and the friend says \"a spot's a spot\" and the Placeholder nods along and something small breaks."))

quest("placeholder_t5_b", 5, "placeholder",
      "The Placeholder has started holding a spot for you. You didn't ask. \"In case you need to step out,\" they say. \"Ever. I've got it. That's what I do.\"",
      ("Accept it gracefully", C, 12,
       "\"Thank you,\" you say, and mean it. They stand a little taller. Someone's spot is being held, and it's yours, and that turns out to matter.",
       "You say \"you don't have to\" and they say \"I know\" and you say \"really\" and now it's a thing, and you made it a thing."),
      ("Refuse to ever step out, so the holding is never needed", S, 15,
       "You don't step out. Not once. Hours. The Placeholder holds a spot that never opens. Neither of you mentions it. It's perfect.",
       "You last a long time and then you need to step out and you do and the Placeholder holds your spot flawlessly and you feel like you've lost a bet."))

quest("vendor_t5_a", 5, "vendor",
      "The Vendor has packed up the table. There's no item today. \"Last day,\" they say. \"Moving on. Different line. I've got one thing left, and it's not for sale.\" They hold out their hand. There's nothing in it.",
      ("Take what's in the hand", I, 15,
       "You take it. It's nothing. You put it in your pocket. The Vendor nods. \"Look after that.\" You will. You're not sure what it is, but you will.",
       "You look at the empty hand and say \"there's nothing there\" and the Vendor closes their hand and says \"not any more\" and leaves."),
      ("Ask them to stay one more day", C, 13,
       "\"One more day,\" they say. They unpack the table. Tomorrow's item is a single sock. You don't buy it. You don't need to.",
       "You ask them to stay and they say \"you never bought anything\" and you didn't, did you. Not once."))

quest("vendor_t5_b", 5, "vendor",
      "The Vendor's item today is a folded piece of paper. \"Fifty dollars. No change. It's a note. I don't know what it says. I don't read the stock.\"",
      ("Deduce where the note came from", I, 16,
       "You examine the fold, the paper, the crease pattern. It's a grocery list. The Vendor knew. \"Forty dollars,\" they say.",
       "You examine it and declare it's something important. It's a grocery list. The Vendor reads it aloud. \"Eggs. Milk.\" Fifty dollars."),
      ("Refuse to look at it, on principle", S, 10,
       "You don't look. You don't ask. The Vendor is impressed. \"Most people look,\" they say. You're not most people. You're in a line.",
       "You refuse, and then your eyes drift, and you see the word \"eggs,\" and the Vendor says \"that's a look\" and charges you a dollar for it."))

quest("auditor_t5_a", 5, "auditor",
      "The Auditor is at the door. Next in. They turn around. \"I've audited nine hundred lines. This is the first one I've reached the front of. I'm scared.\" They hold out the clipboard. \"Sign it. As a witness.\"",
      ("Sign the clipboard", C, 15,
       "You sign. They look at your signature for a long time. Then they turn around, and the door opens, and that's all you know.",
       "You sign, but your hand's shaking too, and it comes out wrong, and they say \"that's not a signature\" and the moment's gone and the door opens anyway."),
      ("Tell them what the audit found", I, 16,
       "You summarise it. Nine hundred lines. One right one. They nod. \"Case closed,\" they say, and step through, and the clipboard stays with you.",
       "You summarise and you get a number wrong and they correct you, at the door, and then they're gone, and the last thing they said was a number."))

quest("auditor_t5_b", 5, "auditor",
      "\"I've worked it out,\" says the Auditor, very calm. \"There's no wrong line. There never was. Every line is the right one for the people in it.\" They look at you. \"Tell me I'm wrong.\"",
      ("Find the flaw in the argument", I, 18,
       "You find it. It's subtle. They listen, and then they smile. \"There it is,\" they say. \"I knew there was one.\" They seem relieved to be wrong again.",
       "You try and you can't. They wait. \"So I'm right,\" they say. They don't look happy about it. Being right was never the point."),
      ("Tell them they're right", C, 14,
       "\"You're right,\" you say. They put the clipboard down on the pavement and leave it there, and stand in line like everyone else. Just a person, now.",
       "You say \"you're right\" too fast and they say \"you didn't even think about it\" and pick the clipboard back up."))

# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------
SHEETS = [
    ("Tuning", ["key", "value", "notes"], TUNING),
    ("XP_Curve", ["level", "xp_required"], XP_CURVE),
    ("Quest_Tiers", ["tier", "unlock_level", "difficulty_min", "difficulty_max", "xp_success"], QUEST_TIERS),
    ("Odds_Bands", ["label", "min_probability"], ODDS_BANDS),
    ("Titles", ["stat", "min_level", "title", "accessory"], TITLES),
    ("Avatar_Parts", ["layer", "option"], AVATAR_PARTS),
    ("Emojis", ["emoji", "ascii", "label"], EMOJIS),
    ("NPCs", ["npc_id", "name", "description", "stat_lean"], NPCS),
    ("NPC_Lines", ["npc_id", "text"], NPC_LINES),
    ("Announcements", ["npc_id", "event", "text"], ANNOUNCEMENTS),
    ("Prompts", ["text"], [(p,) for p in PROMPTS]),
    ("Seed_Messages", ["text", "prompt"], [(m, PROMPTS[pi]) for m, pi in SEED_MESSAGES]),
    ("Blocklist", ["word"], [(w,) for w in BLOCKLIST]),
    ("Quests", ["id", "tier", "npc_id", "setup",
                "a_label", "a_stat", "a_difficulty", "a_success", "a_failure",
                "b_label", "b_stat", "b_difficulty", "b_success", "b_failure"], Q),
]


def write_csvs():
    os.makedirs(CONTENT, exist_ok=True)
    for name, header, rows in SHEETS:
        path = os.path.join(CONTENT, f"{name}.csv")
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(header)
            for r in rows:
                w.writerow(r)
        print(f"wrote {path} ({len(rows)} rows)")


def write_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("README")
    lines = [
        "THE LINE: content + tuning sheet",
        "",
        "Every tab here is loaded by the game. Edit values in place; keep the header rows and tab names as they are.",
        "To pull changes into the game: share the sheet as 'Anyone with the link can view', then run `npm run content:pull` with SHEET_ID set.",
        "",
        "Tabs:",
        "  Tuning         key/value numbers from the design doc. 'notes' explains each.",
        "  XP_Curve       cumulative XP needed to reach each level. Last row is the level cap.",
        "  Quest_Tiers    unlock level, difficulty band and XP payout per tier.",
        "  Odds_Bands     probability floor for each qualitative odds label.",
        "  Titles         title + text accessory by dominant stat and level. Stat NONE is the default when no stat has been raised.",
        "  Avatar_Parts   options for each text-avatar layer. Randomize picks one per layer.",
        "  Emojis         exit emoji menu, with ASCII fallback.",
        "  NPCs           roster.",
        "  NPC_Lines      ambient chat lines per NPC.",
        "  Announcements  templates for wins. {name} and {level} are substituted.",
        "  Prompts        rotating prompt shown above the message box.",
        "  Seed_Messages  hand-written messages seeded into the pool, each with the prompt it answers.",
        "  Blocklist      words that auto-reject in the heuristic moderator.",
        "  Quests         one row per quest. Stats: CHARM, INTELLIGENCE, STRENGTH. Difficulty should sit within the tier's band.",
    ]
    for i, l in enumerate(lines, 1):
        readme.cell(row=i, column=1, value=l)
    readme.column_dimensions["A"].width = 120
    readme["A1"].font = Font(bold=True, size=14)

    head_fill = PatternFill("solid", fgColor="DDDDDD")
    for name, header, rows in SHEETS:
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in range(1, len(header) + 1):
            ws.cell(row=1, column=c).font = Font(bold=True)
            ws.cell(row=1, column=c).fill = head_fill
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        for c, h in enumerate(header, 1):
            width = max([len(str(h))] + [len(str(r[c - 1])) for r in rows]) if rows else len(h)
            ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 70)
            if width > 40:
                for row in range(2, len(rows) + 2):
                    ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    os.makedirs(BUILD, exist_ok=True)
    path = os.path.join(BUILD, "the-line-content.xlsx")
    wb.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    write_csvs()
    if "--no-xlsx" not in sys.argv:
        write_xlsx()
